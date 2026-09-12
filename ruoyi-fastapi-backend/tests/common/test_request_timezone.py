import asyncio
from datetime import datetime, timezone
from io import BytesIO

import httpx
import pytest
from fastapi import FastAPI, Request, status
from openpyxl import load_workbook

from common.annotation.cache_annotation import ApiCache
from common.context import RequestContext
from config.env import AppConfig
from middlewares.context_middleware import add_context_cleanup_middleware
from module_admin.dao.config_dao import ConfigDao
from module_admin.entity.vo.config_vo import ConfigPageQueryModel
from module_admin.entity.vo.user_vo import CurrentUserModel, UserInfoModel
from utils.excel_util import ExcelUtil
from utils.page_util import PageUtil
from utils.time_util import TimezoneUtil


@pytest.fixture
def timezone_app(monkeypatch: pytest.MonkeyPatch) -> FastAPI:
    """用真实中间件、DAO日期边界和Excel导出验证请求时区链路。"""
    monkeypatch.setattr(AppConfig, 'app_timezone', 'Asia/Shanghai')
    app = FastAPI()
    add_context_cleanup_middleware(app)

    async def capture_query(db: object, query: object, *args: object) -> dict:
        return query.compile().params

    monkeypatch.setattr(PageUtil, 'paginate', capture_query)

    @app.get('/range')
    async def date_range(request: Request, preference: str = 'auto') -> dict:
        RequestContext.set_current_user(
            CurrentUserModel(permissions=[], roles=[], user=UserInfoModel(userId=1, timeZone=preference))
        )
        # 交错执行多个请求，以暴露进程全局时区被覆盖的问题。
        await asyncio.sleep(0)
        params = await ConfigDao.get_config_list(
            None, ConfigPageQueryModel(beginTime='2026-03-08', endTime='2026-03-08')
        )
        bounds = sorted(value for value in params.values() if isinstance(value, datetime))
        workbook = load_workbook(
            BytesIO(
                ExcelUtil.export_list2excel([{'at': datetime(2026, 3, 8, 7, 30, tzinfo=timezone.utc)}], {'at': '时间'})
            )
        )
        key = await ApiCache(namespace='test')._build_cache_key(request)
        return {
            'zone': TimezoneUtil.get_request_timezone(),
            'start': TimezoneUtil.format_rfc3339(bounds[0]),
            'end': TimezoneUtil.format_rfc3339(bounds[1]),
            'excelHeader': workbook.active.cell(1, 1).value,
            'excelTime': workbook.active.cell(2, 1).value.isoformat(),
            'cacheKey': key,
            'businessZone': TimezoneUtil.get_app_timezone().key,
        }

    @app.get('/failure')
    async def failure() -> None:
        raise RuntimeError('request failed')

    return app


@pytest.mark.asyncio
async def test_parallel_requests_isolate_dates_exports_and_cache(timezone_app: FastAPI) -> None:
    transport = httpx.ASGITransport(app=timezone_app)
    async with httpx.AsyncClient(transport=transport, base_url='http://test') as client:
        responses = await asyncio.gather(
            *(
                client.get('/range', headers={'X-Timezone': zone})
                for zone in ['America/New_York', 'Asia/Shanghai', 'America/New_York', 'Asia/Shanghai']
            )
        )
    ny, sh, ny_again, sh_again = [response.json() for response in responses]
    assert ny['zone'] == 'America/New_York'
    assert ny['start'] == '2026-03-08T05:00:00.000Z'
    assert ny['end'] == '2026-03-09T04:00:00.000Z'  # 春季切换日只有23小时
    assert ny['excelHeader'] == '时间 (America/New_York)'
    assert ny['excelTime'] == '2026-03-08T03:30:00'
    assert sh['start'] == '2026-03-07T16:00:00.000Z'
    assert sh['end'] == '2026-03-08T16:00:00.000Z'
    assert sh['excelHeader'] == '时间 (Asia/Shanghai)'
    assert sh['excelTime'] == '2026-03-08T15:30:00'
    assert ny['cacheKey'] != sh['cacheKey']
    assert ny == ny_again
    assert sh == sh_again
    assert ny['businessZone'] == sh['businessZone'] == 'Asia/Shanghai'
    assert RequestContext.get_current_timezone() is None


@pytest.mark.asyncio
async def test_request_timezone_fallback_validation_and_exception_cleanup(timezone_app: FastAPI) -> None:
    transport = httpx.ASGITransport(app=timezone_app, raise_app_exceptions=False)
    async with httpx.AsyncClient(transport=transport, base_url='http://test') as client:
        assert (await client.get('/range')).json()['zone'] == 'Asia/Shanghai'
        assert (await client.get('/range', params={'preference': 'America/New_York'})).json()[
            'zone'
        ] == 'America/New_York'
        # 同一账号其他设备修改偏好后，已打开页面上报的时区仍与本次查询一致。
        response = await client.get('/range', params={'preference': 'America/New_York'}, headers={'X-Timezone': 'UTC'})
        assert response.json()['zone'] == 'UTC'
        for value in ['', 'auto', 'UTC+08:00', 'Mars/Olympus', '../UTC']:
            response = await client.get('/range', headers={'X-Timezone': value})
            assert response.status_code == status.HTTP_422_UNPROCESSABLE_CONTENT
            assert 'IANA' in response.json()['msg']
        assert (
            await client.get('/failure', headers={'X-Timezone': 'America/New_York'})
        ).status_code == status.HTTP_500_INTERNAL_SERVER_ERROR
        assert (await client.get('/range')).json()['zone'] == 'Asia/Shanghai'
    assert RequestContext.get_current_timezone() is None


def test_background_exports_can_select_business_timezone_explicitly() -> None:
    token = RequestContext.set_current_timezone('America/New_York')
    try:
        workbook = load_workbook(
            BytesIO(
                ExcelUtil.export_list2excel(
                    [{'at': datetime(2026, 3, 8, 7, 30, tzinfo=timezone.utc)}], {'at': '时间'}, timezone_name='UTC'
                )
            )
        )
        assert workbook.active.cell(1, 1).value == '时间 (UTC)'
        assert workbook.active.cell(2, 1).value == datetime(2026, 3, 8, 7, 30)
    finally:
        RequestContext.reset_current_timezone(token)
