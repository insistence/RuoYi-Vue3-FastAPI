from datetime import datetime, timedelta, timezone
from io import BytesIO
from types import SimpleNamespace

import pytest
import pytest_asyncio
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

from common.context import RequestContext
from config.database import Base
from exceptions.exception import OidcInteractionException
from module_identity.dao.oauth_audit_dao import OAuthAuditDao
from module_identity.entity.do.oauth_audit_do import SysOAuthAuditArchive, SysOAuthAuditLog
from module_identity.entity.vo.oauth_session_vo import AuditPageQueryModel
from module_identity.service.audit_service import AuditService
from utils.oidc_util import OidcUtil

_ADMIN_EVENT_ID = 7
_ADMIN_PAGE_SIZE = 5
_ADMIN_TOTAL = 1
_EXPORT_LIMIT = 5000
_ROLLBACK_COUNT = 2
_SERVICE_UNAVAILABLE = 503


@pytest_asyncio.fixture
async def audit_session() -> AsyncSession:
    """创建只包含审计表的真实 SQLite 会话。"""
    engine = create_async_engine('sqlite+aiosqlite:///:memory:')
    async with engine.begin() as connection:
        await connection.run_sync(
            lambda sync_connection: Base.metadata.create_all(
                sync_connection,
                tables=[SysOAuthAuditLog.__table__, SysOAuthAuditArchive.__table__],
            )
        )
    factory = async_sessionmaker(engine, expire_on_commit=False)
    async with factory() as session:
        yield session
    await engine.dispose()


def test_audit_redacts_nested_protocol_secrets_and_keeps_safe_fields() -> None:
    """嵌套字典和列表中的协议秘密不得落入审计详情。"""
    result = OidcUtil.sanitize_audit_detail(
        {
            'attempt': 2,
            'client': {'name': 'web', 'client_secret': 'hidden'},
            'nested': [
                {'code': 'authorization-code', 'label': 'safe'},
                {'opaque': 'rt1.record-id.secret-value'},
                {'opaque': 'ac1.record-id.secret-value'},
                {'opaque': 'prefix ss1.record-id.secret-value suffix'},
                {'opaque': 'prefix cs1.secret-value-with-enough-entropy suffix'},
                {'pkce_verifier': 'hidden'},
            ],
            'password': 'hidden',
        }
    )
    assert result == {
        'attempt': 2,
        'client': {'name': 'web'},
        'nested': [
            {'label': 'safe'},
            {'opaque': '[REDACTED]'},
            {'opaque': '[REDACTED]'},
            {'opaque': '[REDACTED]'},
            {'opaque': '[REDACTED]'},
            {},
        ],
    }


def test_audit_redacts_opaque_user_agent_but_keeps_plain_identifiers() -> None:
    """凭据外观的 User-Agent 脱敏，普通 Client/Token ID 保持可检索。"""
    event = AuditService.build_event(
        event_type='authorization_code_reused',
        result='failure',
        client_id='cs1',
        token_id='rt1',
        user_agent='Mozilla/5.0 rt1.record-id.secret-value suffix',
    )
    assert event.user_agent == '[REDACTED]'
    assert event.client_id == 'cs1'
    assert event.token_id == 'rt1'
    assert event.risk_level == 'high'


def test_audit_rejects_invalid_event_type_length() -> None:
    """事件类型超过数据库列长度时必须拒绝。"""
    with pytest.raises(ValueError):
        AuditService.build_event(event_type='x' * 65, result='success')


@pytest.mark.asyncio
async def test_audit_uses_field_whitelist_and_caller_commit_boundary(audit_session: AsyncSession) -> None:
    """未知顶层字段被丢弃，服务只 flush 不 commit。"""
    event = await AuditService.record(
        audit_session,
        'invalid_client',
        'failure',
        detail={'safe': 'value', 'access_token': 'hidden'},
        not_a_column='must-not-be-stored',
    )
    assert event.risk_level == 'high'
    assert event.detail == {'safe': 'value'}
    assert not audit_session.new
    loaded = (await audit_session.execute(select(SysOAuthAuditLog))).scalars().first()
    assert loaded is event


@pytest.mark.asyncio
async def test_audit_retention_moves_old_events_to_archive(audit_session: AsyncSession) -> None:
    """超过在线保留期的审计事件必须转入归档表后再删除。"""
    now = datetime.now(timezone.utc)
    await AuditService.record(audit_session, 'old_event', 'success', create_time=now - timedelta(days=181))
    await AuditService.record(audit_session, 'fresh_event', 'success', create_time=now)

    archived = await OAuthAuditDao.archive_before(audit_session, now - timedelta(days=180))
    await audit_session.flush()

    assert archived == 1
    online = (await audit_session.execute(select(SysOAuthAuditLog))).scalars().all()
    archive = (await audit_session.execute(select(SysOAuthAuditArchive))).scalars().all()
    assert [row.event_type for row in online] == ['fresh_event']
    assert [row.event_type for row in archive] == ['old_event']


@pytest.mark.asyncio
async def test_audit_admin_reads_project_safe_models_and_export(monkeypatch: pytest.MonkeyPatch) -> None:
    """管理端分页和导出共用脱敏投影与筛选条件。"""
    row = SimpleNamespace(
        event_id=_ADMIN_EVENT_ID,
        event_type='login',
        result='success',
        risk_level='normal',
        trace_id='trace-1',
        client_id='client-1',
        resource_id=None,
        user_id=3,
        subject_id='subject-1',
        sid='sid-1',
        ip_address='127.0.0.1',
        failure_code=None,
        create_time=datetime.now(timezone.utc),
    )
    calls: list[dict[str, object]] = []

    async def list_admin_page(*args: object, **kwargs: object) -> list[object]:
        calls.append(kwargs)
        return [row]

    async def count_admin(*args: object, **kwargs: object) -> int:
        return 1

    monkeypatch.setattr(OAuthAuditDao, 'list_admin_page', list_admin_page)
    monkeypatch.setattr(OAuthAuditDao, 'count_admin', count_admin)
    query = AuditPageQueryModel(page_num=2, page_size=_ADMIN_PAGE_SIZE, client_id='client-1')

    rows, total = await AuditService.list_admin_page(SimpleNamespace(), query)
    exported = await AuditService.export_admin(SimpleNamespace(), query)

    assert total == _ADMIN_TOTAL
    assert rows[0]['auditId'] == _ADMIN_EVENT_ID
    assert 'detail' not in rows[0]
    workbook = load_workbook(BytesIO(exported))
    headers, values = list(workbook.active.values)
    assert dict(zip(headers, values, strict=True))['auditId'] == _ADMIN_EVENT_ID
    assert 'detail' not in headers
    assert calls[0]['offset'] == _ADMIN_PAGE_SIZE and calls[0]['limit'] == _ADMIN_PAGE_SIZE
    assert calls[1]['offset'] == 0 and calls[1]['limit'] == _EXPORT_LIMIT


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ('timezone_name', 'expected_time'),
    [
        ('Asia/Shanghai', datetime(2026, 8, 28, 18, 30, 0, 123000)),
        ('America/New_York', datetime(2026, 8, 28, 6, 30, 0, 123000)),
    ],
)
async def test_audit_export_serializes_stored_utc_in_request_timezone(
    audit_session: AsyncSession, timezone_name: str, expected_time: datetime
) -> None:
    """真实查询和 Excel 写入链路必须转换 UTC 时刻并注明导出时区。"""
    event = await AuditService.record(
        audit_session,
        'export_timezone',
        'success',
        create_time=datetime(2026, 8, 28, 10, 30, 0, 123000, tzinfo=timezone.utc),
        detail={'private_note': 'not-for-export', 'access_token': 'hidden'},
    )
    await audit_session.commit()
    token = RequestContext.set_current_timezone(timezone_name)
    try:
        exported = await AuditService.export_admin(audit_session, AuditPageQueryModel(event_type='export_timezone'))
    finally:
        RequestContext.reset_current_timezone(token)
    workbook = load_workbook(BytesIO(exported))
    headers, values = list(workbook.active.values)
    exported_row = dict(zip(headers, values, strict=True))
    assert exported_row['auditId'] == event.event_id
    assert exported_row[f'createTime ({timezone_name})'] == expected_time
    assert 'detail' not in headers
    assert 'not-for-export' not in values
    assert 'hidden' not in values


@pytest.mark.asyncio
async def test_interaction_failure_audit_rolls_back_and_wraps_failure(monkeypatch: pytest.MonkeyPatch) -> None:
    """交互失败审计保持双重回滚和统一异常包装语义。"""

    class _Db:
        def __init__(self) -> None:
            self.rollbacks = 0

        async def rollback(self) -> None:
            self.rollbacks += 1

    async def fail(*args: object, **kwargs: object) -> object:
        raise RuntimeError('audit unavailable')

    monkeypatch.setattr(AuditService, 'record_independent', fail)
    db = _Db()
    with pytest.raises(OidcInteractionException) as raised:
        await AuditService.record_interaction_failure(db, 'login_failed', failure_code='invalid_credentials')

    assert raised.value.error == 'server_error'
    assert raised.value.status_code == _SERVICE_UNAVAILABLE
    assert db.rollbacks == _ROLLBACK_COUNT
